"""Build the final report artefacts from the experiment results.

    .venv\\Scripts\\python reports\\build_report.py          # DOCX + PDF (PDF needs Microsoft Word)
    .venv\\Scripts\\python reports\\build_report.py --no-pdf

Outputs
* reports/IEEE_FINAL_REPORT.docx and .pdf
* docs/EXPERIMENTS_AND_RESULTS.md
* reports/figures/*.png (paper figures) and reports/data/<run_id>/ (snapshot of the runs used)
"""
from __future__ import annotations

import os
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

import report_data as rd
from report_data import f

import figures
import paper_figures as pf

ROOT = rd.ROOT
REPORTS = ROOT / "reports"
FIG = REPORTS / "figures"

TITLE = ("A Comparative Study of Performance, Consistency, and Security in NoSQL Databases for Big Data "
         "Applications, with a Proposed Hybrid Adaptive Framework")
# Register numbers are personal data and are not stored in the repository: set REPORT_AUTHOR_REG_NOS
# (comma separated, same order as the names) before building the report.
_NAMES = ["Rohit Kumar", "Yash Krishnadixit", "Dev Pradaan"]
_REGS = [r.strip() for r in os.environ.get("REPORT_AUTHOR_REG_NOS", "").split(",")]
AUTHORS = [(n, _REGS[i] if i < len(_REGS) and _REGS[i] else "[reg. no.]") for i, n in enumerate(_NAMES)]
PLACEHOLDER_DEMO = "[to be added: demo video link]"
PLACEHOLDER_DRIVE = "[to be added: public Google Drive folder link]"


def pct(x, d=1):
    return f"{x:.{d}f}%"


# =============================================================================================
# Text of the paper. Every number comes from `v` (the loaded results).
# =============================================================================================
def paper_text(v) -> dict:
    by, wls = v["e1_by"], v["e1_wls"]
    d = v["e1_delta"]
    ov = v["e1_overhead"]
    rm = v["e2_rm"]
    live = v["e2_live"]
    mig = live["migration"]
    r07_first, r07_last, r07_best, _ = v["e2_reb07"]
    r99_first, r99_last, _, _ = v["e2_reb099"]
    s3 = v["e3_summary"]
    rules = s3["Static rules"]
    sel = v["e3_sel_mean"]
    sel_c = ", ".join(f"{c:g}" for c in v["e3_sel_c"])
    s4 = v["e4_summary"]
    ad, q, al, one = (s4["HANF adaptive"], s4["static QUORUM/QUORUM"], s4["static ALL/ALL"], s4["static ONE/ONE"])
    orc = s4["adaptive (oracle)"]
    sd = {r["read_delay_ms"]: r for r in v["e4_sens_delay"] if r["strategy"] == "static ONE/ONE"}
    env = v["e1_env"]
    sweep = v["e3_sweep_fn"]
    n_w = len(wls)
    bvs = v["e1_batched_beats_sync"]
    avs = v["e1_async_beats_sync"]
    stub_r = ov.get("full_pipeline_read_no_io_us")
    stub_u = ov.get("full_pipeline_update_no_io_us")
    p50_gap = [by[(w, "hanf")]["p50_us"] - by[(w, "direct")]["p50_us"] for w in wls]
    n_runs = len(v["e1_all_runs"])
    robust_sync = sum(r["sync_has_highest_p99_everywhere"] for r in v["e1_all_runs"])
    robust_p50 = sum(r["hanf_p50_above_direct_everywhere"] for r in v["e1_all_runs"])

    T = {}
    T["abstract"] = (
        "NoSQL engines are usually chosen and tuned with benchmarks that look at speed, consistency or "
        "security one at a time, although the three interact in every deployment. This paper completes the "
        "Hybrid Adaptive NoSQL Framework (HANF), a middleware layer that takes four decisions for each "
        "request: the engine (Redis, MongoDB or Cassandra) chosen by an explainable workload classifier, the "
        "Redis shard chosen by consistent hashing with virtual nodes, the read/write consistency level chosen "
        "by a Monte-Carlo latency and staleness predictor, and an anomaly verdict produced by an Isolation "
        "Forest over per-client request windows. The modules exchange signals through one telemetry bus and "
        "are exposed through a REST API and a web dashboard. Four experiments were re-run on the finished "
        "implementation. On live Redis, the complete HANF path cost "
        f"{-d['hanf'][1]:.0f}-{-d['hanf'][0]:.0f}% of single-client throughput relative to direct access, and "
        f"inline per-window model scoring raised p99 latency to {v['e1_p99']['hanf_sync_ml'][0]/1000:.1f}-"
        f"{v['e1_p99']['hanf_sync_ml'][1]/1000:.1f} ms, whereas batched scoring cut the per-window model cost by a "
        f"factor of about {v['e3_single_ms']/v['e3_batch_per_window_ms']:.0f}. Adding a fifth shard moved "
        f"{mig['remapped_pct']:.1f}% of {mig['total_before']:,} live keys with no lost or mismatched value. On "
        f"held-out synthetic traffic the Isolation Forest reached F1 = {sel['f1']:.3f} against {rules['f1']:.3f} "
        "for static rules, but missed slow brute-force attacks. In simulation, adaptive consistency kept stale "
        f"reads at {ad['stale_pct']:.1f}% while lowering mean latency by {-ad['mean_vs_quorum_pct']:.1f}% "
        "relative to static QUORUM. The paper separates live measurements from simulated and synthetic results "
        "and lists the remaining limitations.")
    T["keywords"] = ("NoSQL, polyglot persistence, consistent hashing, tunable consistency, anomaly detection, "
                     "Isolation Forest, Redis, MongoDB, Cassandra, middleware")

    T["lit"] = [
        ("A. Background", [
            "Key-value, document and wide-column stores trade strict consistency and rich queries for "
            "horizontal scale and flexible schemas. Replicated stores such as Dynamo and Cassandra let the client "
            "choose how many replicas must acknowledge a write (W) and answer a read (R); when W + R exceeds the "
            "replica count N the read and write sets intersect, otherwise a read may return an old value [13], "
            "[14]. Brewer's CAP discussion explains why this choice cannot be avoided under partitions [11], and "
            "the probabilistically bounded staleness (PBS) model shows that partial quorums are often fresh in "
            "practice and that staleness can be predicted from latency distributions [12]. Data placement in "
            "these systems usually relies on consistent hashing, which moves only about 1/(N+1) of the keys when "
            "a node joins [8].",
        ]),
        ("B. Related work", [
            "Abbasi et al. benchmarked Redis, MongoDB, Cassandra, DynamoDB Local and Couchbase with YCSB [7] on "
            "datasets of 100 thousand to 10 million records and derived selection criteria; Redis was fastest at "
            "small scale but lost about a third of its throughput at large scale, while Cassandra improved as data "
            "grew [6]. Bodra and Khairnar compared Redis, Aerospike and Dragonfly under identical YCSB workloads "
            "and found that the multi-threaded engines overtake Redis on some multi-threaded workloads [5]. "
            "Ferreira et al. measured MongoDB, Redis and Cassandra deployed across cloud regions and showed that "
            "geographic placement strongly affects the time to reach consistency [1]; a companion study reported "
            "throughput losses that depend on the engine and operation when stronger levels are configured [2]. "
            "Solat's review of sharding concludes that consistent hashing balances load and limits data movement "
            "better than static hash or range partitioning, but proposes no new benchmark [4]. On security, Ron "
            "et al. showed that operator injection (for example $ne or $where in a JSON query) is a practical "
            "attack on document stores [15], and machine-learning anomaly detection over access logs has been "
            "proposed as a stronger alternative to static rules [3]. Isolation Forest [9] is well suited to this "
            "setting because it needs no attack labels and scores samples cheaply.",
        ]),
        ("C. Gaps and problem statement", [
            "The Review I survey of these studies found three recurring gaps: each study evaluates one dimension "
            "(performance, consistency or security) in isolation; experiments use a single hardware or cloud "
            "configuration; and consistency levels, sharding and security controls are configured statically "
            "instead of reacting to the workload. The problem addressed here is therefore to build and evaluate "
            "one middleware layer that makes routing, placement, consistency and security decisions together, "
            "shares their signals, and reports which results are measured on live engines and which are "
            "simulated.",
        ]),
        ("D. Contributions", [
            "Relative to the Review II prototype (about 61% complete by self-assessment, Redis only), this work "
            "(i) completes the router with an explainable classifier and live MongoDB and Cassandra adapters; "
            "(ii) makes shard migration verified, idempotent and failure-tolerant; (iii) moves ML scoring off the "
            "request path with batched and asynchronous modes and separates model selection from testing; "
            "(iv) connects the consistency controller to the live request path; (v) adds an event-based "
            "telemetry bus, a REST API, a dashboard, 92 automated tests and reproducible experiment scripts.",
        ]),
    ]

    T["arch"] = [
        "Fig. 1 shows the architecture. Client applications call a REST API built with FastAPI. Each request "
        "passes through a fixed pipeline: validation, workload classification and engine choice, shard lookup "
        "for the key-value engine, consistency-level lookup, execution on an engine adapter, and publication of "
        "a telemetry event. The security monitor is a subscriber of the telemetry bus, so it observes every "
        "request without sitting on the request path. The sharding manager and the consistency controller also "
        "publish their decisions on the bus, and the dashboard derives all of its counters from bus events.",
        "Three engines run in Docker: five independent Redis 7.0.15 instances on ports 7001-7005 (four active "
        "shards and one spare, as in Review II), MongoDB 7.0 and Cassandra 4.1. Every engine is reached through "
        "the same adapter interface, which also reports health; an engine is shown as connected only after a "
        "real write, read and delete have succeeded. Redis is the dependable path for the demonstration: if "
        "MongoDB or Cassandra is unavailable, the router falls back to Redis and flags the decision in the "
        "response. Table I summarises module status.",
    ]

    T["impl"] = [
        ("A. Workload-aware polyglot router", [
            "The classifier scores three workload classes with transparent weighted rules held in a YAML file: an "
            "explicit client hint (+10), operations that only one engine supports (scan and login for key-value, "
            "field queries for documents, +6), known key prefixes such as product:, order: or sensor: (+3), a "
            "scalar value (+2 key-value), a nested object or array (+2.5 document), an object with a clustering "
            "field ts (+3.5 wide-column) and a small tie-breaker for key-value. The highest total wins and its "
            "share of the total is reported as confidence, together with every rule that fired. The document "
            "adapter rejects field names that begin with $ or contain a dot and accepts only scalar equality "
            "filters, so an operator payload never reaches MongoDB as an operator [15]. The wide-column adapter "
            "stores rows under a partition key with the timestamp as clustering key and applies the controller's "
            "level to every CQL statement.",
        ]),
        ("B. Adaptive sharding manager", [
            "Keys and virtual-node tokens are hashed with the first 64 bits of MD5 and a key is owned by the next "
            "token clockwise. When a shard is added or removed the ring switches immediately, writes go to the new "
            "owner and reads fall back to the previous owner until migration ends. For each key whose owner "
            "changes, the manager copies the serialized value with DUMP and RESTORE, compares the destination copy "
            "byte by byte with the source, and only then deletes the source copy. A destination value that "
            "already differs was written after the switch, so it is kept. Failed keys remain on the source, the "
            "report is marked partial and a retry completes the work; total key counts before and after are "
            "reconciled in every report. Load-based rebalancing re-weights the virtual nodes with the Review II "
            "rule (2), using per-shard request counts from the bus:",
        ]),
        ("C. ML security monitor", [
            "Each run of 100 consecutive requests from one client forms a window described by six features: "
            "request rate, key diversity (distinct keys over requests), mean request plus response bytes, failure "
            "ratio, the share of requests matching an injection pattern ($ne, $where, UNION, SELECT, DROP, OR "
            "1=1, comment markers) and the share of SCAN requests. Features are z-scored with the training "
            "statistics and scored by an Isolation Forest with 100 trees trained only on normal windows; a window "
            "is flagged when its anomaly score exceeds the threshold implied by the contamination parameter. "
            "Static rules (injection > 0.2, failure > 0.5, scan > 0.5, rate > 150 requests/s) run in parallel as a "
            "baseline. Windows can be scored inline, in batches of 20, or by a background worker (default). Each "
            "verdict lists the features that deviate most from training, and only an analyst's triage marks an "
            "event as a confirmed attack. Training and test traffic come from a seeded generator of normal, "
            "injection, exfiltration and brute-force windows, 30% of them low-intensity variants; the same "
            "generator can execute its requests through the live pipeline for the demonstration.",
        ]),
        ("D. Adaptive consistency controller", [
            "The controller models N = 3 replicas in regions A, B and C with a writer in A and a reader in B "
            "(mean RTTs 2/30/110 ms and 30/2/80 ms). Each sampled RTT is the base value times a congestion factor "
            "on remote links times LogNormal(0, 0.25) jitter. The write latency is the W-th smallest writer RTT; a "
            "read starts Exp(15 ms) after the acknowledgement, contacts the R closest replicas and is stale if none "
            "of them has applied the write when the request arrives, following the PBS reasoning [12]. For seven "
            "(W, R) profiles the controller estimates mean latency and stale probability from Monte-Carlo samples "
            "and applies rule (3); in the request path the chosen level is applied to Cassandra statements and "
            "recorded as advisory for Redis and MongoDB.",
        ]),
        ("E. Telemetry, interfaces and quality assurance", [
            "The bus is an in-process publish/subscribe component with timestamped events, a 5,000-event "
            "retention window and isolation of subscriber failures. The REST API offers 33 endpoints "
            "(documented with OpenAPI) for requests, seeding, shards, security, consistency, background benchmark "
            "jobs and exports. The dashboard has eight pages and calls only real endpoints. The test suite has 92 "
            "tests: unit tests for every module, API tests on in-process Redis emulation, experiment tests, and "
            "live integration tests against the Redis, MongoDB and Cassandra containers, including a forced "
            "mid-migration failure. A 15-step smoke test drives the whole demonstration flow over HTTP.",
        ]),
    ]
    T["eq"] = [
        "owner(k) = node of min{ t : h(t) >= h(k) } on the ring",
        "v_i <- max(16, round(v_i (1 - 0.5 (L_i / mean(L) - 1))))",
        "(W, R)* = argmin { E[latency](W, R) : P(stale | W, R) <= SLO }, else (ALL, ALL)",
    ]

    T["setup"] = (
        f"All measurements were taken on one laptop: {env['cpu']}, {env['logical_cpus']} logical CPUs, "
        f"{env['memory_gb']} GB RAM, {env['os'].split(' (')[0]}, Python {env['python']}, scikit-learn "
        f"{env['libraries']['scikit-learn']}, Redis {env['engines'].get('redis_version')} in Docker Desktop. Runs "
        "that use Redis write to a separate logical database that is flushed afterwards. E1 was run from the "
        "command line with the API server, MongoDB and Cassandra stopped. Table II lists the evidence type of "
        "each experiment; configuration, environment metadata, raw tables and graphs of every run are stored "
        "with the code.")

    # ---------------- E1
    e1 = [
        f"E1 reproduces the YCSB core workloads A (50% reads, 50% updates), B (95/5) and C (reads only) with "
        f"scrambled Zipfian keys (theta 0.99) over {v['e1_cfg']['records']:,} records of {v['e1_cfg']['value_bytes']} bytes, "
        f"{v['e1_cfg']['operations']:,} operations per run and {v['e1_cfg']['repetitions']} repetitions, issued "
        "sequentially by one client thread. Table III and Fig. 2 report the results. Direct Redis reached "
        f"{f(v['e1_direct_tp'][0], 0)}-{f(v['e1_direct_tp'][1], 0)} operations/s with a median latency of about "
        f"{v['e1_p50']['direct'][0]:.0f}-{v['e1_p50']['direct'][1]:.0f} us. These absolute values are much lower "
        "than native Redis because every request crosses Docker Desktop's virtualised network, so the "
        "configurations should be compared with each other rather than with published figures. "
        f"E1 was executed {len(v['e1_all_runs'])} times with the default profile; background desktop load made "
        f"some runs very noisy (worst direct-Redis coefficient of variation "
        f"{max(r['direct_cv_max'] for r in v['e1_all_runs']):.2f}), so the run reported here is the most recent "
        f"one whose direct-Redis coefficient of variation is below {v['e1_rule']:.2f}, a rule fixed before the "
        "last run was inspected; all runs are listed with the published results.",
        f"The full HANF path without ML scoring reduced throughput by {-d['hanf'][1]:.1f}% to {-d['hanf'][0]:.1f}% "
        f"and added {min(p50_gap):.0f}-{max(p50_gap):.0f} us to the median latency. To locate this cost, the "
        "same pipeline was timed with the Redis connections replaced by an in-memory stub: it needed "
        + (f"{stub_r:.0f} us per read and {stub_u:.0f} us per update, of which classification took "
           f"{ov['router_classify_route_us']:.1f} us, the ring lookup {ov['hash_ring_lookup_us']:.1f} us and the "
           f"telemetry event {ov['telemetry_publish_us']:.1f} us"
           + (" (measured in a separate E1 run). " if v["e1_overhead_source"] != v["runs"]["e1"]["run_id"] else ". ")
           if stub_r else "about 30 us per request. ")
        + "The rest of the gap is therefore not computation in HANF; it coincides with spreading requests over "
        "four Redis connections through Docker's port-forwarding proxy instead of one, which could not be "
        "isolated further on this machine.",
        f"Inline (synchronous) scoring of every completed window lowered throughput by {-d['hanf_sync_ml'][1]:.1f}%-"
        f"{-d['hanf_sync_ml'][0]:.1f}% relative to direct access and raised p99 latency to "
        f"{v['e1_p99']['hanf_sync_ml'][0]/1000:.1f}-{v['e1_p99']['hanf_sync_ml'][1]/1000:.1f} ms, because the "
        "request that completes a window waits for a model call. In the reported run, batched scoring kept more "
        f"throughput than synchronous scoring in {bvs} of {n_w} workloads (losses of {-d['hanf_batched_ml'][1]:.1f}%-"
        f"{-d['hanf_batched_ml'][0]:.1f}%) and asynchronous scoring in {avs} of {n_w} "
        f"({-d['hanf_async_ml'][1]:.1f}%-{-d['hanf_async_ml'][0]:.1f}%), with p99 latency of "
        f"{v['e1_p99']['hanf_batched_ml'][0]/1000:.1f}-{v['e1_p99']['hanf_batched_ml'][1]/1000:.1f} ms and "
        f"{v['e1_p99']['hanf_async_ml'][0]/1000:.1f}-{v['e1_p99']['hanf_async_ml'][1]/1000:.1f} ms. The "
        f"asynchronous worker had at most {v['e1_backlog']:.0f} windows waiting when a run ended, so it kept up. "
        "Across all E1 runs, however, the throughput ordering of the three scoring modes changed from run to "
        "run, so we do not rank them by throughput. Two findings held in every run: synchronous scoring had the "
        f"highest p99 latency of all configurations in every workload ({robust_sync} of {n_runs} runs), and the "
        f"HANF median latency was above that of direct access ({robust_p50} of {n_runs} runs). The cost of a model "
        f"call explains the tail effect: about {v['e3_single_ms']:.1f} ms whether it scores one window or twenty "
        f"({v['e3_batch_per_window_ms']:.2f} ms per window in a batch of 20, measured in E3).",
    ]
    T["e1"] = e1

    # ---------------- E2
    T["e2"] = [
        f"Adding one node to {v['e2_cfg']['remap_keys']:,} keys (Table IV, Fig. 3a; mean of "
        f"{v['e2_cfg']['ring_seeds']} ring seeds) moved {rm[('modulo', '4->5')]['remapped_pct_mean']:.1f}% and "
        f"{rm[('modulo', '10->11')]['remapped_pct_mean']:.1f}% of keys with modulo hashing, but only "
        f"{rm[('consistent v=128', '4->5')]['remapped_pct_mean']:.1f}% and "
        f"{rm[('consistent v=128', '10->11')]['remapped_pct_mean']:.1f}% with consistent hashing and 128 virtual "
        "nodes, close to the ideal 20.0% and 9.1%. More virtual nodes reduced the spread across seeds (SD "
        f"{rm[('consistent v=16', '4->5')]['remapped_pct_sd']:.1f} points at 16 virtual nodes versus "
        f"{rm[('consistent v=256', '4->5')]['remapped_pct_sd']:.1f} at 256) and lowered the initial imbalance from "
        f"{rm[('consistent v=16', '4->5')]['imbalance_mean']:.2f} to {rm[('consistent v=256', '4->5')]['imbalance_mean']:.2f} "
        "(maximum over mean key count).",
        f"Under Zipf-0.7 request skew, the rebalancing rule lowered the hottest node's load from "
        f"{r07_first['max_over_mean']:.3f} to {r07_best['max_over_mean']:.3f} times the mean (coefficient of "
        f"variation {r07_first['cv']:.3f} to {min(r['cv'] for r in v['e2_reb07'][3]):.3f}) within a few rounds, "
        "each adjustment moving only a small share of keys (Fig. 3b). Under Zipf-0.99 skew a single key carried "
        f"{r99_first['hottest_key_share_pct']:.1f}% of all requests, and the imbalance stayed between "
        f"{min(r['max_over_mean'] for r in v['e2_reb099'][3]):.2f} and {max(r['max_over_mean'] for r in v['e2_reb099'][3]):.2f}: "
        "re-weighting virtual nodes cannot split one hot key, which needs replication or caching instead.",
        f"In the live test, {live['keys_loaded']:,} keys were loaded into four Redis shards and a fifth shard was "
        f"added. The manager moved {mig['moved']:,} keys ({mig['remapped_pct']:.1f}%) in {mig['elapsed_s']:.2f} s, "
        f"verified each one at the destination, and reported {mig['failed']} failures. An independent check read "
        f"{live['sample_verified']['checked']:,} sampled values from their new owners "
        f"({live['sample_verified']['mismatches']} mismatches), found {live['keys_not_on_ring_owner']} keys stored "
        f"on the wrong shard, and confirmed {mig['total_before']:,} keys before and {mig['total_after']:,} after.",
    ]

    # ---------------- E3
    sub = v["e3_cfg"]
    T["e3"] = [
        f"The detector was trained on {sub['train_windows']:,} normal windows. The contamination value was "
        f"chosen on a separate validation set ({sub['val_normal']:,} normal and {sub['val_per_attack']} windows of "
        f"each attack) and all metrics were measured on an independent test set ({sub['test_normal']:,} normal and "
        f"{sub['test_per_attack']} windows of each attack, 30% subtle); the protocol was repeated for "
        f"{len(sub['seeds'])} seeds. Review II had compared contamination values on the test set itself; the "
        "split removes that bias.",
        f"Validation selected contamination {sel_c} in every seed. On the test sets the Isolation Forest reached "
        f"precision {sel['precision']:.3f}, recall {sel['recall']:.3f}, F1 {sel['f1']:.3f} +/- {v['e3_sel_sd']['f1']:.3f} and a "
        f"false-positive rate of {sel['fpr']*100:.2f}%, and detected {sel['subtle_recall']*100:.1f}% of the subtle "
        f"attacks. Static rules were almost never wrong (precision {rules['precision']:.3f}) but found only "
        f"{rules['recall']*100:.1f}% of attacks (F1 {rules['f1']:.3f}), because every subtle variant stays below "
        "their thresholds (Table V, Fig. 4a). Larger contamination values raised the false-positive rate to "
        f"{s3['Isolation Forest (c=0.05)']['fpr']*100:.1f}% at 0.05 without improving recall.",
        "The evasion sweep (Fig. 4b) inserts a growing share of attack requests into otherwise normal windows. "
        f"Injection was detected in {sweep('injection', 0.04)*100:.0f}% of windows at 4% attack requests and in "
        f"all windows from 8%. Exfiltration needed about 16% scan requests for {sweep('exfiltration', 0.16)*100:.0f}% "
        f"recall. Brute force spread over normal-rate traffic was the weakest case: even at 30% failed logins "
        f"only {sweep('bruteforce', 0.3)*100:.0f}% of windows were flagged, so slow password guessing needs a "
        "per-account failure counter in addition to the window model. All of these results come from synthetic "
        "traffic and do not show accuracy on real logs.",
    ]

    # ---------------- E4
    T["e4"] = [
        "E4 runs the 16-phase congestion timeline of Review II (multipliers 1-6, 4,000 operations per phase, "
        f"{len(v['e4_cfg']['seeds'])} seeds) under static ONE/ONE, QUORUM/QUORUM and ALL/ALL and under the "
        "adaptive controller, which re-selects at the start of each phase from the congestion it observed in the "
        f"previous phase (Table VI, Fig. 5). ONE/ONE was fastest ({one['mean_ms']:.1f} ms) but "
        f"{one['stale_pct']:.1f}% of its reads were stale. QUORUM/QUORUM and ALL/ALL were never stale but averaged "
        f"{q['mean_ms']:.1f} ms and {al['mean_ms']:.1f} ms. The adaptive controller had {ad['stale_pct']:.1f}% stale "
        f"reads, a mean of {ad['mean_ms']:.1f} ms ({-ad['mean_vs_quorum_pct']:.1f}% below QUORUM and "
        f"{-ad['mean_vs_all_pct']:.1f}% below ALL) and a p95 {-ad['p95_vs_quorum_pct']:.1f}% below QUORUM.",
        "The controller alternated between ONE/QUORUM and QUORUM/ONE. Neither satisfies W + R > N, so their "
        "freshness is not guaranteed by quorum intersection; it holds here because the replica nearest the "
        "writer is always among the two replicas nearest the reader. The simulation discovers this property of "
        "the layout; with another placement the predicted staleness, and hence the choice, would change. For the "
        "same reason the result did not depend on the SLO between 0.1% and 5%, and the oracle that knows the "
        f"current congestion was no better ({orc['mean_ms']:.1f} ms): both remaining profiles have the same "
        "expected latency, so the one-phase lag cost nothing in this scenario. The ONE/ONE staleness depends "
        f"strongly on the read delay: {sd[5]['stale_pct']:.1f}%, {sd[15]['stale_pct']:.1f}% and "
        f"{sd[50]['stale_pct']:.1f}% for mean delays of 5, 15 and 50 ms.",
        "Compared with Review II, the ordering of strategies and the size of the saving are similar (Review II "
        "reported 50.4% below QUORUM), but absolute values differ: Review II's QUORUM read latency of 31.9 ms "
        "is not reachable under its own stated model, in which a QUORUM read from region B must wait for the "
        "30 ms link to region A multiplied by an average congestion of 2.4, so the earlier prototype probably "
        "differed in an undocumented detail. All E4 numbers are simulated.",
    ]

    T["threats"] = (
        "The main threats to validity are the single shared laptop with Docker networking, a single-threaded "
        "Python load generator instead of the official YCSB client, synthetic security traffic with attack "
        "families designed by the authors, a consistency model with assumed latencies, and single-node MongoDB "
        "and Cassandra deployments. Fixed seeds, repeated runs, stored environment metadata and the separation "
        "of validation and test data reduce but do not remove these threats.")

    T["conclusion"] = [
        "HANF is now a working middleware layer in which routing, sharding, consistency selection and security "
        "monitoring run in one request pipeline over Redis, MongoDB and Cassandra and share one telemetry bus. "
        "The experiments support three conclusions. Consistent hashing with verified migration can grow a live "
        "cluster while moving close to the minimum number of keys and losing none. An unsupervised window model "
        "finds attacks that static thresholds miss at a false-positive rate below one percent on synthetic data, "
        "but it must be scored in batches or asynchronously to protect tail latency, and it does not catch slow "
        "brute force. In simulation, choosing the read/write profile from predicted latency and staleness halves "
        "the latency of static QUORUM without stale reads, but the benefit depends on replica placement.",
        "Future work should deploy a multi-node, multi-region Cassandra cluster to measure the controller on "
        "real replication, evaluate the security monitor on a labelled real query log and against adaptive "
        "attackers, add per-account counters for slow brute force, replicate hot keys, and repeat E1 with the "
        "official YCSB client on dedicated hardware.",
    ]
    return T


# =============================================================================================
# Tables (rows built from results)
# =============================================================================================
def tables(v) -> dict:
    by, wls = v["e1_by"], v["e1_wls"]
    names = {"direct": "Direct Redis", "hanf": "HANF routing", "hanf_sync_ml": "+ sync ML",
             "hanf_batched_ml": "+ batched ML", "hanf_async_ml": "+ async ML"}
    e1 = []
    for w in wls:
        for c in ("direct", "hanf", "hanf_sync_ml", "hanf_batched_ml", "hanf_async_ml"):
            if (w, c) not in by:
                continue
            r = by[(w, c)]
            e1.append([w, names[c], f"{r['throughput_mean']:,.0f} ± {r['throughput_sd']:,.0f}",
                       "–" if c == "direct" else f"{r['delta_vs_direct_pct']:+.1f}%",
                       f"{r['p50_us']:.0f}", f"{r['p99_us']:,.0f}"])
    rm = v["e2_rm"]
    e2 = []
    for s in ("modulo", "consistent v=16", "consistent v=64", "consistent v=128", "consistent v=256"):
        a, b = rm[(s, "4->5")], rm[(s, "10->11")]
        e2.append([s.replace("consistent", "Consistent,").replace("modulo", "Modulo (hash mod N)"),
                   f"{a['remapped_pct_mean']:.1f} ± {a['remapped_pct_sd']:.1f}",
                   f"{b['remapped_pct_mean']:.1f} ± {b['remapped_pct_sd']:.1f}", f"{a['imbalance_mean']:.2f}"])
    e2.append(["Ideal 1/(N+1)", "20.0", "9.1", "1.00"])
    s3 = v["e3_summary"]
    e3 = []
    for m, r in s3.items():
        e3.append([m.replace("Isolation Forest", "IF"), f"{r['precision']:.3f}", f"{r['recall']:.3f}",
                   f"{r['f1']:.3f}", f"{r['fpr']*100:.2f}", f"{r['recall_injection']*100:.0f}",
                   f"{r['recall_exfiltration']*100:.0f}", f"{r['recall_bruteforce']*100:.0f}",
                   f"{r['subtle_recall']*100:.0f}", str(r["times_selected"])])
    s4 = v["e4_summary"]
    e4 = []
    for n in ("static ONE/ONE", "static QUORUM/QUORUM", "static ALL/ALL", "HANF adaptive", "adaptive (oracle)"):
        r = s4[n]
        e4.append([n.replace("static ", "Static "), f"{r['write_ms']:.1f}", f"{r['read_ms']:.1f}",
                   f"{r['mean_ms']:.1f}", f"{r['p95_ms']:.1f}", f"{r['stale_pct']:.2f}"])
    status = [
        ["Polyglot Router", "35% (KV path only)", "Done: classifier, Redis/MongoDB/Cassandra adapters, fallback", "live"],
        ["Adaptive Sharding", "90%", "Done: verified idempotent migration, retry, load re-weighting", "live + algorithmic"],
        ["ML Security Monitor", "70% (shadow mode)", "Done: async/batched scoring, validation split, triage", "synthetic"],
        ["Consistency Controller", "60% (offline sim.)", "Done: in request path; applied on Cassandra", "simulation"],
        ["Telemetry Bus", "50%", "Done: event bus, retention, derived metrics", "live"],
    ]
    evidence = [["E1", "Throughput/latency, YCSB A/B/C", "Live Redis (Docker)"],
                ["E2", "Remapping, rebalancing, live migration", "Algorithmic + live Redis"],
                ["E3", "IF vs static rules, evasion sweep", "Synthetic traffic"],
                ["E4", "Static vs adaptive consistency", "Monte-Carlo simulation"]]
    return {"e1": e1, "e2": e2, "e3": e3, "e4": e4, "status": status, "evidence": evidence}


REFERENCES = [
    'S. Ferreira, J. Mendonça, B. Nogueira, W. Tiengo, and E. Andrade, "Impacts of data consistency levels in cloud-based NoSQL for data-intensive applications," J. Cloud Comput., vol. 13, no. 1, art. 158, 2024, doi: 10.1186/s13677-024-00716-7.',
    'S. Ferreira, J. Mendonça, and E. Andrade, "Experimental performance analysis of data consistency levels in NoSQL databases," Softw.: Pract. Exper., vol. 55, no. 6, pp. 1059–1070, 2025, doi: 10.1002/spe.3412.',
    'A. Maamari and M. Nasar, "Enhancing database security in NoSQL systems using machine learning-based anomaly detection," 2025 [venue to be confirmed].',
    'S. Solat, "Sharding distributed databases: A critical review," arXiv:2404.04384, 2024.',
    'D. Bodra and S. Khairnar, "Comparative performance analysis of modern NoSQL data technologies: Redis, Aerospike, and Dragonfly," J. Res. Innov. Technol., vol. 4, no. 2, pp. 193–200, 2025; arXiv:2510.08863.',
    'M. Abbasi, P. Váz, J. Silva, F. Cardoso, F. Sá, and P. Martins, "Multi-database NoSQL performance analysis: Empirical selection framework for cloud-native applications," Procedia Comput. Sci., vol. 280, pp. 745–752, 2026, doi: 10.1016/j.procs.2026.04.094.',
    'B. F. Cooper, A. Silberstein, E. Tam, R. Ramakrishnan, and R. Sears, "Benchmarking cloud serving systems with YCSB," in Proc. 1st ACM Symp. Cloud Comput. (SoCC), 2010, pp. 143–154, doi: 10.1145/1807128.1807152.',
    'D. Karger, E. Lehman, T. Leighton, R. Panigrahy, M. Levine, and D. Lewin, "Consistent hashing and random trees: Distributed caching protocols for relieving hot spots on the World Wide Web," in Proc. 29th ACM Symp. Theory Comput. (STOC), 1997, pp. 654–663, doi: 10.1145/258533.258660.',
    'F. T. Liu, K. M. Ting, and Z.-H. Zhou, "Isolation forest," in Proc. 8th IEEE Int. Conf. Data Mining (ICDM), 2008, pp. 413–422, doi: 10.1109/ICDM.2008.17.',
    'F. Pedregosa et al., "Scikit-learn: Machine learning in Python," J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.',
    'E. Brewer, "CAP twelve years later: How the \'rules\' have changed," Computer, vol. 45, no. 2, pp. 23–29, 2012, doi: 10.1109/MC.2012.37.',
    'P. Bailis, S. Venkataraman, M. J. Franklin, J. M. Hellerstein, and I. Stoica, "Probabilistically bounded staleness for practical partial quorums," Proc. VLDB Endow., vol. 5, no. 8, pp. 776–787, 2012, doi: 10.14778/2212351.2212359.',
    'G. DeCandia et al., "Dynamo: Amazon\'s highly available key-value store," ACM SIGOPS Oper. Syst. Rev., vol. 41, no. 6, pp. 205–220, 2007, doi: 10.1145/1323293.1294281.',
    'A. Lakshman and P. Malik, "Cassandra: A decentralized structured storage system," ACM SIGOPS Oper. Syst. Rev., vol. 44, no. 2, pp. 35–40, 2010, doi: 10.1145/1773912.1773922.',
    'A. Ron, A. Shulman-Peleg, and A. Puzanov, "Analysis and mitigation of NoSQL injections," IEEE Secur. Privacy, vol. 14, no. 2, pp. 30–39, 2016, doi: 10.1109/MSP.2016.36.',
]


# =============================================================================================
# DOCX writer
# =============================================================================================
def build_docx(v, T, tabs, figs, out: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.9)
    sec.left_margin = sec.right_margin = Inches(0.66)

    def cols(section, n):
        sp = section._sectPr
        c = sp.find(qn("w:cols"))
        if c is None:
            c = OxmlElement("w:cols")
            sp.append(c)
        c.set(qn("w:num"), str(n))
        c.set(qn("w:space"), "288")

    def rich(p, text, size=None, italic=False, bold=False):
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not part:
                continue
            b, i = bold, italic
            if part.startswith("**"):
                part, b = part[2:-2], True
            elif part.startswith("*"):
                part, i = part[1:-1], True
            r = p.add_run(part)
            r.bold, r.italic = b, i
            if size:
                r.font.size = Pt(size)
        return p

    def para(text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, after=0, italic=False, bold=False):
        p = doc.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.17)
        p.paragraph_format.space_after = Pt(after)
        return rich(p, text, size, italic, bold)

    roman = iter(["I", "II", "III", "IV", "V", "VI", "VII"])

    def heading(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"{next(roman)}. {text}")
        r.font.small_caps = True
        r.font.size = Pt(10)

    def subheading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.italic = True

    def equation(text, n):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.italic = True
        r.font.size = Pt(9)
        p.add_run(f"   ({n})").font.size = Pt(9)

    tcount = iter(["I", "II", "III", "IV", "V", "VI", "VII"])

    def table(title, header, rows, widths=None, size=7.5):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"TABLE {next(tcount)}\n")
        r.font.size = Pt(8)
        r = p.add_run(title)
        r.font.size = Pt(8)
        r.font.small_caps = True
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate([header] + rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
                rr = cp.add_run(str(val))
                rr.font.size = Pt(size)
                rr.bold = i == 0
        for row in t.rows[:-1]:  # keep the whole table in one column
            for cell in row.cells:
                for cp_ in cell.paragraphs:
                    cp_.paragraph_format.keep_with_next = True
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
        if widths:
            for j, w in enumerate(widths):
                for row in t.rows:
                    row.cells[j].width = Inches(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    fcount = iter(range(1, 20))

    def figure(path, caption, width):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        c.paragraph_format.space_after = Pt(6)
        r = c.add_run(f"Fig. {next(fcount)}. ")
        r.font.size = Pt(8)
        r = c.add_run(caption)
        r.font.size = Pt(8)

    # ------------------------------------------------------------ cover page (DA-3 template)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Digital Assignment 3 – Implementation Report")
    r.font.size = Pt(13)
    r.bold = True
    cover = [("Course Code and Name", "BCSE406L – NoSQL Databases"), ("Slot", "D1"),
             ("Register Number of the Student", ", ".join(a[1] for a in AUTHORS)),
             ("Name of the Student", ", ".join(a[0] for a in AUTHORS)),
             ("Title of the Paper", TITLE), ("DEMO video link", PLACEHOLDER_DEMO),
             ("Course Faculty", "Manjula V")]
    t = doc.add_table(rows=len(cover), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, val) in enumerate(cover):
        for j, txt in enumerate((k, val)):
            c = t.cell(i, j)
            c.text = ""
            rr = c.paragraphs[0].add_run(txt)
            rr.font.size = Pt(11)
            rr.bold = j == 0
            c.width = Inches(2.4 if j == 0 else 4.4)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    rich(p, "Sections as required by the DA-3 template: Abstract; I. Literature Review and Problem Statement; "
            "II. Proposed System Architecture; III. Proposed System and Algorithm Implementation; IV. Results "
            "and Discussion; V. Graphs, Plots and Tables; VI. Conclusion; References; Appendix (Drive link with "
            "code, demo video and user manual).", size=10, italic=True)

    # ------------------------------------------------------------ title block (one column)
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    cols(s, 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(TITLE)
    r.font.size = Pt(20)
    at = doc.add_table(rows=1, cols=3)
    for j, (name, reg) in enumerate(AUTHORS):
        c = at.cell(0, j)
        c.text = ""
        cp = c.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run(name).font.size = Pt(11)
        for line in (f"Reg. No. {reg}", "Vellore Institute of Technology", "Chennai, India"):
            rr = cp.add_run("\n" + line)
            rr.font.size = Pt(9)
            rr.italic = not line.startswith("Reg")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------ body (two columns)
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    cols(s, 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Abstract—")
    r.bold, r.italic = True, True
    r.font.size = Pt(9)
    r = p.add_run(T["abstract"])
    r.bold = True
    r.font.size = Pt(9)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("Index Terms—")
    r.bold, r.italic = True, True
    r.font.size = Pt(9)
    r = p.add_run(T["keywords"])
    r.bold = True
    r.font.size = Pt(9)

    heading("Literature Review and Problem Statement")
    for title, ps in T["lit"]:
        subheading(title)
        for x in ps:
            para(x)

    heading("Proposed System Architecture")
    para(T["arch"][0])
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    cols(s, 1)
    figure(figs["arch"], "HANF architecture. The request pipeline calls the router, the sharding manager and the "
           "consistency controller; the security monitor subscribes to request events on the shared telemetry "
           "bus. Redis, MongoDB and Cassandra are reached through one adapter interface.", 6.2)
    s = doc.add_section(WD_SECTION.CONTINUOUS)
    cols(s, 2)
    para(T["arch"][1])
    table("Module status: Review II versus this work", ["Module", "Review II", "This work", "Evidence"],
          tabs["status"], widths=[0.85, 0.7, 1.25, 0.6], size=7)

    heading("Proposed System and Algorithm Implementation")
    for i, (title, ps) in enumerate(T["impl"]):
        subheading(title)
        for x in ps:
            para(x)
        if i == 1:
            equation(T["eq"][0], 1)
            equation(T["eq"][1], 2)
            para("where v_i is the number of virtual nodes of shard i and L_i its observed request load; "
                 "the floor of 16 keeps every shard on the ring.", indent=False)
        if i == 3:
            equation(T["eq"][2], 3)

    heading("Results and Discussion")
    subheading("A. Experimental setup")
    para(T["setup"])
    table("Experiments and type of evidence", ["ID", "Experiment", "Evidence"], tabs["evidence"],
          widths=[0.35, 1.7, 1.3])
    subheading("B. E1: Throughput and latency")
    for x in T["e1"]:
        para(x)
    subheading("C. E2: Adaptive sharding")
    for x in T["e2"]:
        para(x)
    subheading("D. E3: ML security monitor")
    for x in T["e3"]:
        para(x)
    subheading("E. E4: Adaptive consistency (simulation)")
    for x in T["e4"]:
        para(x)
    subheading("F. Threats to validity")
    para(T["threats"])

    heading("Graphs, Plots and Tables")
    para("Tables III–VI and Figs. 2–5 are generated by the experiment scripts from the stored result files; "
         "± values are standard deviations over repetitions or seeds.")
    table("E1 – Throughput (ops/s) and latency (µs), live Redis, 3 runs",
          ["WL", "Configuration", "Throughput", "Δ vs direct", "p50", "p99"], tabs["e1"],
          widths=[0.3, 0.95, 0.9, 0.55, 0.35, 0.45], size=7)
    figure(figs["e1"], "E1 on live Redis: (a) throughput and (b) p99 latency of direct access and the HANF "
           "configurations for YCSB-style workloads A, B and C.", 3.35)
    table("E2 – Keys remapped (%) when one node is added (100,000 keys, 5 ring seeds)",
          ["Scheme", "4→5", "10→11", "Imbalance (4 nodes)"], tabs["e2"], widths=[1.3, 0.75, 0.75, 0.6], size=7)
    figure(figs["e2"], "E2: (a) share of keys moved when a node is added, (b) hottest-node load over "
           "rebalancing rounds for two levels of request skew.", 3.35)
    table("E3 – Detection on held-out synthetic test windows (mean of 3 seeds)",
          ["Method", "Prec.", "Rec.", "F1", "FPR %", "Inj. %", "Exf. %", "BF %", "Subtle %", "Sel."],
          tabs["e3"], size=6.5)
    figure(figs["e3"], "E3 on synthetic traffic: (a) precision, recall and F1 of static rules and Isolation "
           "Forest for several contamination values; (b) recall as attacks are diluted into normal windows.", 3.35)
    table("E4 – Static and adaptive consistency (simulation, mean of 3 seeds)",
          ["Strategy", "Write ms", "Read ms", "Mean ms", "p95 ms", "Stale %"], tabs["e4"],
          widths=[1.15, 0.45, 0.45, 0.5, 0.5, 0.45], size=7)
    figure(figs["e4"], "E4 (simulation): (a) mean latency and (b) stale reads per strategy; (c) the adaptive "
           "controller's observed and predicted latency in each phase of the congestion timeline.", 3.35)

    heading("Conclusion")
    for x in T["conclusion"]:
        para(x)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("References")
    r.font.small_caps = True
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r = p.add_run(f"[{i}]\t{ref}")
        r.font.size = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Appendix")
    r.font.small_caps = True
    runs = v["runs"]
    appendix = [
        f"**Drive link (code with access, demo video, user manual):** {PLACEHOLDER_DRIVE}",
        f"**Demo video:** {PLACEHOLDER_DEMO}",
        "**Code:** backend/ (FastAPI application and tests), frontend/ (dashboard), scripts/ (setup, start, "
        "smoke test, experiments), configs/, docker-compose.yml. Start with README.md.",
        "**User manual:** docs/USER_MANUAL.md; run guide docs/HOW_TO_RUN.md; demo script docs/DEMO_SCRIPT.md.",
        "**Runs used in this paper:** " + "; ".join(f"{k.upper()} {r['run_id']}" for k, r in runs.items())
        + ". Their files are in reports/data/.",
        "**Reproduce:** scripts\\setup.ps1 -Dev; docker compose up -d; scripts\\run_experiments.ps1 -Experiment "
        "all -Profile default; python reports\\build_report.py.",
    ]
    for x in appendix:
        para(x, indent=False, size=8.5, after=2)
    doc.save(out)


def export_pdf(docx: Path, pdf: Path) -> dict:
    ps = (
        "$ErrorActionPreference='Stop';"
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false;"
        f"$d = $w.Documents.Open('{docx}');"
        "$d.Repaginate();"
        "$pages = $d.ComputeStatistics(2); $words = $d.ComputeStatistics(0);"
        f"$d.SaveAs([ref]'{pdf}', [ref]17); $d.Close([ref]0); $w.Quit();"
        "Write-Output \"$pages $words\""
    )
    out = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=300)
    if out.returncode != 0:
        raise RuntimeError(out.stderr.strip()[-800:])
    pages, words = out.stdout.strip().split()[-2:]
    return {"pages_total_incl_cover": int(pages), "words_total": int(words)}


def body_word_count(T) -> int:
    texts = [T["abstract"], *T["arch"], T["setup"], *T["e1"], *T["e2"], *T["e3"], *T["e4"], T["threats"],
             *T["conclusion"]]
    for _, ps in T["lit"] + T["impl"]:
        texts += ps
    return sum(len(t.split()) for t in texts)


def main() -> int:
    v = rd.load()
    rd.snapshot(v["runs"], [r["run_id"] for r in v["e1_all_runs"]])
    FIG.mkdir(parents=True, exist_ok=True)
    figs = {"arch": figures.architecture(FIG / "fig_architecture.png"),
            "e1": pf.e1(v, FIG), "e2": pf.e2(v, FIG), "e3": pf.e3(v, FIG), "e4": pf.e4(v, FIG)}
    figures.pipeline(FIG / "fig_pipeline.png")
    T = paper_text(v)
    tabs = tables(v)
    import results_md
    results_md.write(v, T, tabs, ROOT / "docs" / "EXPERIMENTS_AND_RESULTS.md")
    docx = REPORTS / "IEEE_FINAL_REPORT.docx"
    build_docx(v, T, tabs, figs, docx)
    print("wrote", docx, f"(body text ~{body_word_count(T)} words)")
    if "--no-pdf" not in sys.argv:
        info = export_pdf(docx, REPORTS / "IEEE_FINAL_REPORT.pdf")
        print("wrote", REPORTS / "IEEE_FINAL_REPORT.pdf", info)
    return 0


if __name__ == "__main__":
    sys.exit(main())
